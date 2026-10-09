"""Create the public-safe survey dataset from the private original workbook."""

from __future__ import annotations

import argparse
import csv
from pathlib import Path

from openpyxl import load_workbook


TIME_OPTIONS = ["Early Morning", "Morning", "Afternoon", "Evening", "Late Night"]
EVENT_OPTIONS = ["Sporting Events", "Conventions", "Tourism", "General Travel", "Concerts/music"]
RETENTION_OPTIONS = [
    "Increase rates",
    "More promotions/opportunities",
    "Ability for customers to request a specific driver",
    "Improved safety features",
    "Local support office",
    "Eliminate promotion/opportunity zone",
    "Vehicle maintenance tips and tricks",
]

PUBLIC_FIELDS = [
    "respondent_id",
    "time_of_day",
    "driver_tenure",
    "weekly_hours",
    "hourly_income_before_expenses",
    "high_volume_events",
    "local_tourist_ratio",
    "retention_priorities",
    "income_expectation",
]


def selected_options(value: object, allowed: list[str]) -> str:
    """Keep recognized survey choices and discard attached free-text comments."""
    raw = str(value or "").replace("\xa0", " ")
    selected = []
    for option in allowed:
        if option.casefold() in raw.casefold():
            selected.append(option)
    return ";".join(selected)


def clean_expectation(value: object) -> str:
    labels = {
        "Exceeds expectations": "Exceeded expectations",
        "I make what I expected I would make": "Met expectations",
        "Close but not quite": "Close, but below expectations",
        "Not even close": "Not close to expectations",
    }
    return labels.get(str(value or "").strip(), "")


def build_dataset(source: Path, destination: Path) -> int:
    worksheet = load_workbook(source, data_only=True, read_only=True)["Raw data"]
    headers = [cell.value for cell in worksheet[1]]
    rows = []

    for values in worksheet.iter_rows(min_row=2, values_only=True):
        record = dict(zip(headers, values))
        expectation = clean_expectation(record.get("Does your income from Uber meet your expectations?"))
        if not expectation:
            continue

        rows.append(
            {
                "respondent_id": f"R{len(rows) + 1:03d}",
                "time_of_day": selected_options(
                    record.get("What time of day do you work (Select all that apply)?"), TIME_OPTIONS
                ),
                "driver_tenure": str(record.get("How long have you been driving for Uber?") or "").strip(),
                "weekly_hours": str(
                    record.get("On average, how many hours do you drive for Uber on a weekly basis?") or ""
                ).strip(),
                "hourly_income_before_expenses": str(
                    record.get("What is your average hourly income before expenses?") or ""
                ).strip(),
                "high_volume_events": selected_options(
                    record.get("In your experience, what events draw the HIGHEST volume of customers in your market?"),
                    EVENT_OPTIONS,
                ),
                "local_tourist_ratio": str(
                    record.get(
                        "When carrying passengers what is the ratio of locals to tourists? "
                        "(E.G., 25% locals / 75% tourists)"
                    )
                    or ""
                ).strip(),
                "retention_priorities": selected_options(
                    record.get("What do you think Uber can do to increase driver retention (Check all that apply)?"),
                    RETENTION_OPTIONS,
                ),
                "income_expectation": expectation,
            }
        )

    destination.parent.mkdir(parents=True, exist_ok=True)
    with destination.open("w", newline="", encoding="utf-8") as handle:
        writer = csv.DictWriter(handle, fieldnames=PUBLIC_FIELDS)
        writer.writeheader()
        writer.writerows(rows)
    return len(rows)


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("source", type=Path, help="Path to the private original survey workbook")
    parser.add_argument(
        "--output",
        type=Path,
        default=Path("data/uber_driver_survey_anonymized.csv"),
    )
    args = parser.parse_args()
    count = build_dataset(args.source, args.output)
    print(f"Wrote {count} anonymized responses to {args.output}")


if __name__ == "__main__":
    main()

